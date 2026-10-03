import openpyxl,re,json
import sys
F=sys.argv[1]
wb=openpyxl.load_workbook(F,data_only=True)
def num(x):
    return x if isinstance(x,(int,float)) and not isinstance(x,bool) else None
def clean(s): return re.sub(r'\s+',' ',str(s)).strip()
cat=[]
def add(c,id,name,price,**x):
    cat.append(dict(cat=c,id=id,name=name,price=round(price),**x))
def slug(s): return re.sub(r'[^a-z0-9]+','-',s.lower()).strip('-')

# ---------- CPU
CPUSPEC={ # key: (tdp,score)
 'i3-12100F':(58,45),'i3-12100':(60,46),'i3-13100':(60,50),'i3-13100F':(60,50),'i3-14100F':(60,52),'i3-14100':(60,52),
 'i5-12400F':(65,58),'i5-12400':(65,58),'i5-13400F':(65,68),'i5-13400':(65,68),'i5-13500':(65,72),'i5-14400F':(65,70),'i5-14400':(65,70),
 'i5-14600K':(125,90),'i5-14600KF':(125,90),'i7-12700F':(65,78),'i7-12700':(65,78),'i7-12700K':(125,82),'i7-12700КF':(125,82),
 'i7-13700':(65,90),'i7-13700K':(150,96),'i7-13700KF':(150,96),'i7-14700':(65,95),'i7-14700K':(150,100),'i7-14700KF':(150,100),
 'i9-14900K':(200,125),'i9-14900KF':(200,125),
 'Ultra 5-225F':(65,72),'Ultra 5-225':(65,72),'Ultra 5-245K':(160,90),'Ultra 5-245KF':(160,90),'Ultra 7-265K':(200,105),'Ultra 7-265KF':(200,105),'Ultra 9-285K':(250,125),
 '7500F':(65,72),'7600X':(105,80),'7700':(65,85),'7700X':(105,90),'7800X3D':(120,100),
 '5500':(65,52),'5600':(65,60),'5600X':(65,62),'5700X':(65,70),'5500x3D':(105,68),
 '9600X':(65,84),'9700X':(65,92),'9800X3D':(120,108),'9850X3D':(120,112),'9900X':(120,110),'9900X3D':(120,115),'9950X3D':(170,140)}
ws=wb['CPU, DDR, HDD']
for r in ws.iter_rows(min_row=2,max_row=59,values_only=True):
    n,p=r[0],num(r[2])
    if not n or p is None: continue
    n=clean(n); socket=None
    if 'LGA1700' in n: socket='LGA1700'
    elif 'LGA 1851' in n: socket='LGA1851'
    elif 'AM5' in n: socket='AM5'
    elif 'AM4' in n: socket='AM4'
    if not socket: continue
    if n.startswith('Intel'):
        m=re.search(r'Intel-Core (Ultra \d-\w+|i\d) ?-? ?(\w+)',n)
        key=(m.group(1)+('-'+m.group(2) if not m.group(1).startswith('Ultra') else '')) if False else None
        m=re.match(r'Intel-Core (Ultra \d)-(\w+)|Intel-Core (i\d) - (\w+)',n)
        key=(m.group(1)+'-'+m.group(2)) if m.group(1) else (m.group(3)+'-'+m.group(4))
        label='Core '+key.replace('Ultra ','Ultra ')
        igpu=not key.rstrip('K').endswith('F') and not key.endswith('F')
    else:
        m=re.search(r'Ryzen™ (\d) [A-Za-z ]+? (\d{4}\w*) ?-',n); key=m.group(2); label='Ryzen '+m.group(1)+' '+key
        igpu='No GPU' not in n
    if key not in CPUSPEC: print('NO CPU SPEC',key); continue
    tdp,sc=CPUSPEC[key]
    add('cpu','cpu-'+slug(key),label,p,socket=socket,tdp=tdp,score=sc,igpu=igpu)

# ---------- RAM
ram={}
for r in ws.iter_rows(min_row=61,max_row=98,values_only=True):
    n,p=r[0],num(r[2])
    if not n or p is None or 'SODIMM' in n.upper(): continue
    n=clean(n); t=re.search(r'DDR([45])',n); g=re.search(r'DDR[45]\s+(\d+)GB',n)
    if not t or not g: continue
    typ='DDR'+t.group(1); gb=int(g.group(1)); k=(typ,gb)
    if k not in ram or p<ram[k][0]: ram[k]=(p,n)
for (typ,gb),(p,n) in sorted(ram.items()):
    add('ram',f'ram-{typ.lower()}-{gb}',f'{typ} {gb}GB',p,type=typ,gb=gb)

# ---------- SSD
ssd={}
for r in ws.iter_rows(min_row=124,max_row=174,values_only=True):
    n,p=r[0],num(r[2])
    if not n or p is None: continue
    n=clean(n)
    if n.startswith('Ext') or n.startswith('HDD'): continue
    m=re.search(r'(\d+)\s?(TB|GB)',n)
    if not m: continue
    gb=int(m.group(1))*(1000 if m.group(2)=='TB' else 1)
    if gb<240: continue
    bucket={240:256,256:256,500:512,512:512,1000:1000,2000:2000,4000:4000}.get(gb,gb)
    nv='NVMe' in n
    k=(bucket,nv)
    if k not in ssd or p<ssd[k][0]: ssd[k]=(p,n)
for (gb,nv),(p,n) in sorted(ssd.items()):
    lab=f"SSD {gb//1000}TB" if gb>=1000 else f"SSD {gb}GB"
    add('ssd',f"ssd-{gb}-{'nvme' if nv else 'sata'}",lab+(' NVMe' if nv else ' SATA'),p,gb=gb)

# ---------- MB
ws=wb['MB, GPU']
CHIP=r'(H610|B760|Z790|B860|Z890|B550|B450|X570|A620|B650E?|B840|B850|X870E?)(AM|M|i|I)?(?![A-Za-z0-9])'
SOCK={'H610':'LGA1700','B760':'LGA1700','Z790':'LGA1700','B860':'LGA1851','Z890':'LGA1851','B550':'AM4','B450':'AM4','X570':'AM4'}
mbs={}
for r in ws.iter_rows(min_row=2,max_row=206,values_only=True):
    n,p=r[0],num(r[2])
    if not n or p is None: continue
    n=clean(n)
    m=re.search(CHIP,n)
    if not m or 'TRX40' in n: continue
    chip=m.group(1); suf=m.group(2) or ''
    sock=SOCK.get(chip,'AM5')
    t=re.search(r'DDR([45])',n)
    if sock=='AM4': ddr='DDR4'
    elif t: ddr='DDR'+t.group(1)
    else: continue
    if re.search(r'B850i|B850I|B650E-I|B850-I|B650I',n): form='ITX'
    elif suf in('M','AM'): form='mATX'
    else: form='ATX'
    k=(sock,chip,form,ddr)
    if k not in mbs or p<mbs[k][0]: mbs[k]=(p,n)
for (sock,chip,form,ddr),(p,n) in sorted(mbs.items()):
    nm=re.sub(r'^MB\s+','',n); nm=re.sub(r'\s*\(.*?\)','',nm); nm=re.sub(r'\s+(LGA ?\d+|DDR[45]|AMD AM[45])','',nm).strip()
    add('mb',f'mb-{slug(chip+form+ddr+sock)}',f'{nm} ({form})',p,socket=sock,ram=ddr,form=form)

# ---------- GPU
GPUSPEC=[ # regex, label, tdp, len, score
 (r'GT730','GT 730',25,150,6),(r'RX550','RX 550',50,170,8),(r'GT1050Ti','GT 1050 Ti 4GB',75,170,18),(r'RX580','RX 580 8GB',185,245,32),
 (r'GTX1650','GTX 1650 4GB',75,200,25),(r'GTX1660 Super','GTX 1660 Super 6GB',125,230,40),(r'GTX2060 Super','RTX 2060 Super 8GB',175,250,46),
 (r'GTX3050|RTX3050','RTX 3050 8GB',130,240,36),(r'RTX3060','RTX 3060 12GB',170,242,55),(r'RTX5050','RTX 5050 8GB',130,220,48),
 (r'RTX5060Ti','RTX 5060 Ti',180,245,79),(r'RTX5060','RTX 5060 8GB',145,240,66),(r'RTX5070Ti','RTX 5070 Ti 16GB',300,330,125),
 (r'RTX5070','RTX 5070 12GB',250,300,100),(r'RTX5080','RTX 5080 16GB',360,330,150),(r'RTX5090','RTX 5090 32GB',575,360,230),
 (r'RX9060XT','RX 9060 XT 16GB',160,250,78),(r'RX9070XT','RX 9070 XT 16GB',304,330,130),(r'RX9070','RX 9070 16GB',220,320,115)]
gp={}
for r in ws.iter_rows(min_row=209,max_row=359,values_only=True):
    n,p=r[0],num(r[2])
    if not n or p is None: continue
    n=clean(n)
    for rx,lab,tdp,ln,sc in GPUSPEC:
        if re.search(rx,n):
            if lab=='RTX 5060 Ti':
                lab='RTX 5060 Ti '+('16GB' if '16GB' in n else '8GB')
            if lab not in gp or p<gp[lab][0]: gp[lab]=(p,tdp,ln,sc)
            break
    else: print('NO GPU SPEC',n)
for lab,(p,tdp,ln,sc) in sorted(gp.items()):
    add('gpu','gpu-'+slug(lab),lab,p,tdp=tdp,len=ln,score=sc+(1 if '16GB' in lab and 'Ti' in lab else 0))

# ---------- cases / psu / coolers from all sheets
SHEETS=['Cooler, PowerSupply, UPS','Gigabyte','MSI','Montech','Deepcool GamerStorm','TRYX','ID Cooling','Gamdias','Lian Li','FSP','Thermaltake','Jonsbo','PowerCase','CoolerMaster','Redragon']
items=[]
for sn in SHEETS:
    ws=wb[sn]; sec=''
    for r in ws.iter_rows(values_only=True):
        cells=[(j,clean(x)) for j,x in enumerate(r[:7]) if isinstance(x,str) and clean(x)]
        nums=[(j,x) for j,x in enumerate(r[:7]) if num(x) is not None]
        if cells and not nums and len(cells)==1 and r[0]: sec=cells[0][1]; continue
        if not cells or not nums: continue
        namecell=max(cells,key=lambda c:len(c[1]))
        if len(namecell[1])<4: continue
        pr=[x for j,x in nums if j>namecell[0]]
        if not pr: continue
        label=' '.join(c[1] for c in cells if c[0]<namecell[0])
        items.append((sn,sec,label,namecell[1],pr[0]))
print(len(items),'candidate extra rows')
def kind(sn,sec,label,name):
    s=(sec+' '+label+' '+name)
    if re.search(r'UPS|DVD|HDD|SSD|кресл|Chair|Наушник|Headset|мыш|Mouse|клавиат|Keyboard|коврик|Колонк|акуст|микрофон|Brick|Брикет|Рамка|контролер|RGB контр|вентилятор для корпуса|Fan\b',s,re.I) and not re.search(r'Корпус|Case',name): 
        if not re.search(r'cooler|кулер|куллер|водян',s,re.I): return None
    if re.search(r'Корпус|Case|кейс|Кейсы',s,re.I) and not re.search(r'кулер|cooler',sec,re.I):
        if re.search(r'БП|PSU|с водяным|водян|СЖО|\+',name+label) and not re.search(r'Блоки питания',sec): return None
        if re.search(r'Блок|PSU|Блоки',sec): pass
        else: return 'case'
    if re.search(r'Блок питан|Блоки питан|PSU|Power Supply|БЛОКИ ПИТАНИЯ|Блок пит',s,re.I) : return 'psu'
    if re.search(r'cooler|кулер|куллер|CORELIQUID|FROZN|SL360|SL240|FX360|DX360|DASHFLOW|PINKFLOW|BLITZ|SE-\d|Liquid|охлажд|AIO|PANORAMA|SPHERE|Hyper',s,re.I): return 'cooler'
    return None
K={'case':[], 'psu':[], 'cooler':[]}
for sn,sec,label,name,p in items:
    k=kind(sn,sec,label,name)
    if k: K[k].append((sn,sec,label,name,p))
for k,v in K.items(): print(k,len(v))
json.dump(K,open('extra_raw.json','w'),ensure_ascii=False,indent=0)
json.dump(cat,open('cat_main.json','w'),ensure_ascii=False)
print(len(cat),'main items')
